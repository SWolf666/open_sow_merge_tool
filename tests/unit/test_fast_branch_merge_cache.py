from __future__ import annotations

from openpyxl import Workbook

from sow_merge_tool import branch_submit as branch_submit_module
from sow_merge_tool import fast_branch_merge as fast_merge


def _book(path, value):
    book = Workbook()
    book.active["A1"] = value
    book.save(path)
    book.close()


def test_workbook_index_signature_contains_content_hash_and_revision(tmp_path):
    path = tmp_path / "cache.xlsx"
    _book(path, "before")
    fast_merge.cache_clear()
    first = fast_merge.load_workbook_index(str(path), revision=7)
    same = fast_merge.load_workbook_index(str(path), revision=7)
    assert same is first
    assert len(first.signature[1]) == 64
    assert first.signature[2] == 7

    _book(path, "after")
    changed = fast_merge.load_workbook_index(str(path), revision=7)
    assert changed is not first
    assert changed.signature[1] != first.signature[1]

    revision_changed = fast_merge.load_workbook_index(str(path), revision=8)
    assert revision_changed is not changed
    assert revision_changed.signature[1] == changed.signature[1]
    assert revision_changed.signature[2] == 8


def test_branch_engine_source_delta_cache_is_hash_and_revision_scoped(tmp_path, monkeypatch):
    plan = branch_submit_module.FilePlan(
        relative_path="Data.xlsx",
        source_before="before.xlsx",
        source_after="after.xlsx",
        source_before_hash="before-hash",
        source_after_hash="after-hash",
        source_revision=17,
    )
    engine = branch_submit_module.BranchSubmitEngine(str(tmp_path))
    calls = []

    def analyze(before, after, **kwargs):
        calls.append((before, after, kwargs))
        return object()

    monkeypatch.setattr(branch_submit_module, "fast_analyze_source", analyze)
    first = engine._source_delta_for_plan(plan)
    assert engine._source_delta_for_plan(plan) is first
    assert len(calls) == 1

    plan.source_revision = 18
    second = engine._source_delta_for_plan(plan)
    assert second is not first
    assert len(calls) == 2
    assert calls[-1][2] == {"before_revision": 18, "source_revision": 18}

    plan.source_after_hash = "new-after-hash"
    third = engine._source_delta_for_plan(plan)
    assert third is not second
    assert len(calls) == 3
