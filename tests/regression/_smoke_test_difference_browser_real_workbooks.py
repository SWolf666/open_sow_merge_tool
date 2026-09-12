"""Headless production-path checks for update91 difference projections."""

from __future__ import annotations

import hashlib
import shutil
import tempfile
import time
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from sow_merge_tool import legacy_core as smt
from sow_merge_tool.vertical_layout import VerticalLayout


def _write_book(path: Path, rows: list[str], *, extra_sheet: bool = False) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    for row, value in enumerate(rows, 1):
        sheet.cell(row=row, column=1, value=value)
    if extra_sheet:
        cached = book.create_sheet("Cached")
        for row, value in enumerate(rows, 1):
            cached.cell(row=row, column=1, value=value)
    book.save(path)
    book.close()


def _write_column_book(path: Path, headers: list[str], values: list[object]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Columns"
    for column, value in enumerate(headers, 1):
        sheet.cell(row=1, column=column, value=value)
    for column, value in enumerate(values, 1):
        sheet.cell(row=2, column=column, value=value)
    book.save(path)
    book.close()


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _workbook_fingerprint(workbook) -> tuple:
    if workbook is None:
        return ()
    result = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(tuple(row))
        result.append((sheet_name, tuple(rows)))
    return tuple(result)


def _path_hashes(*paths: str | Path) -> dict[str, str]:
    return {
        str(path): _digest(Path(path))
        for path in dict.fromkeys(str(path) for path in paths)
        if path and Path(path).is_file()
    }


def _wait_for_items(app: smt.SowMergeApp, sheet: str, *, timeout: float = 12.0):
    deadline = time.monotonic() + timeout
    items = ()
    while time.monotonic() < deadline:
        app.root.update()
        items = tuple(item for item in app.get_difference_items(5000) if item.sheet == sheet)
        if any("共同删除" in item.summary for item in items) and any(
            "共同新增" in item.summary for item in items
        ):
            return items
        time.sleep(0.01)
    return items


def _wait_for_view(app: smt.SowMergeApp, sheet: str, *, timeout: float = 12.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.root.update()
        view = app.sheet_views.get(sheet)
        if view is not None and view._data_ready and view.column_projection is not None:
            return view
        time.sleep(0.01)
    return app.sheet_views.get(sheet)


def _wait_for_edit_ready(app: smt.SowMergeApp, *, timeout: float = 30.0) -> bool:
    """Wait through the non-blocking editable-workbook preload gate.

    The production UI deliberately keeps browsing/navigation responsive while
    editable workbooks load.  Mutation assertions must therefore distinguish a
    safe temporary rejection from a ready, executable edit path instead of
    calling the synchronous fallback as part of the smoke test.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.root.update()
        if app._edit_workbooks_ready():
            return True
        time.sleep(0.01)
    return app._edit_workbooks_ready()


def _assert_common_row_projection(items) -> None:
    assert any(item.kind is smt.DifferenceKind.DELETED and "共同删除" in item.summary for item in items), items
    assert any(item.kind is smt.DifferenceKind.ADDED and "共同新增" in item.summary for item in items), items


def main() -> None:
    root_dir = Path(tempfile.mkdtemp(prefix="sow-difference-browser-real-"))
    values_base = [f"row-{index}" for index in range(1, 11)]
    values_changed = ["row-1-new", values_base[1], "row-3-new"] + values_base[3:4] + values_base[5:] + ["row-11"]
    base = root_dir / "base.xlsx"
    mine = root_dir / "mine.xlsx"
    theirs = root_dir / "theirs.xlsx"
    _write_book(base, values_base, extra_sheet=True)
    _write_book(mine, values_changed, extra_sheet=True)
    _write_book(theirs, values_changed, extra_sheet=True)
    app = None
    cross_app = None
    columns_app = None
    two_way = None
    try:
        app = smt.SowMergeApp(str(mine), str(theirs), merge_mode=True, base_path=str(base))
        loaded_items = _wait_for_items(app, "Data")
        _assert_common_row_projection(loaded_items)
        loaded_view = app.sheet_views.get("Data")
        assert loaded_view is not None and loaded_view._data_ready
        assert getattr(app._wb_a_val, "read_only", False)
        assert getattr(app._wb_b_val, "read_only", False)
        assert app.ws_a_val("Data")["A1"].value == "row-1-new"
        assert app.ws_b_val("Data")["A1"].value == "row-1-new"
        if not app._edit_workbooks_ready():
            # A click during preload is a non-blocking, safe rejection.  The
            # next click after the gate is ready must be executable.
            loaded_view.selected_pair_idx = 0
            assert loaded_view._copy_selected_row("MINE2A") is False
        assert _wait_for_edit_ready(app)
        # A normal three-way entry (without the legacy conflict-mode flag)
        # must expose an executable retain command that only marks the real
        # DifferenceItem and never changes any workbook bytes.
        retain_item = next(item for item in loaded_items if item.column)
        retain_pair = int(retain_item.id.split(":pair:", 1)[1].split(":", 1)[0])
        loaded_view.selected_pair_idx = retain_pair
        loaded_view._last_selected_line = loaded_view.row_to_line.get(retain_pair, 1)
        loaded_view._selected_difference_item_id = retain_item.id
        assert getattr(app._wb_a_val, "read_only", False)
        assert getattr(app._wb_b_val, "read_only", False)
        retain_paths = (mine, theirs, app.file_a, app.file_b)
        retain_disk_before = _path_hashes(*retain_paths)
        retain_memory_before = (
            _workbook_fingerprint(app._wb_a_val),
            _workbook_fingerprint(app._wb_b_val),
            _workbook_fingerprint(app._wb_a_edit),
            _workbook_fingerprint(app._wb_b_edit),
        )
        assert loaded_view._copy_selected_row("MINE2A")
        assert not getattr(app._wb_a_val, "read_only", True)
        assert not getattr(app._wb_b_val, "read_only", True)
        assert hasattr(app.ws_a_val("Data"), "_cells")
        assert _path_hashes(*retain_paths) == retain_disk_before
        assert (
            _workbook_fingerprint(app._wb_a_val),
            _workbook_fingerprint(app._wb_b_val),
            _workbook_fingerprint(app._wb_a_edit),
            _workbook_fingerprint(app._wb_b_edit),
        ) == retain_memory_before
        assert retain_item.id in app._difference_processed
        assert loaded_view._undo_last_action() is None
        assert _path_hashes(*retain_paths) == retain_disk_before
        assert (
            _workbook_fingerprint(app._wb_a_val),
            _workbook_fingerprint(app._wb_b_val),
            _workbook_fingerprint(app._wb_a_edit),
            _workbook_fingerprint(app._wb_b_edit),
        ) == retain_memory_before
        assert retain_item.id not in app._difference_processed

        # A browser selection can point at row 1 while the user then selects
        # row 2 in the main grid.  If pair lookup misses, the stale browser ID
        # must not win over the current grid pair.
        current_item = next(
            item
            for item in loaded_items
            if item.column
            and ":pair:" in item.id
            and item.id.split(":pair:", 1)[1].split(":", 1)[0]
            != str(retain_pair)
        )
        current_pair = int(current_item.id.split(":pair:", 1)[1].split(":", 1)[0])
        loaded_view.selected_pair_idx = current_pair
        loaded_view._last_selected_line = loaded_view.row_to_line.get(current_pair, 1)
        loaded_view._selected_difference_item_id = retain_item.id
        stale_pair_marker = app.mark_difference_pair_processed
        stale_pair_miss = lambda sheet, pair_idx, columns=None, processed=True: stale_pair_marker(
            sheet, int(pair_idx) + 1000000, columns, processed
        )
        stale_before = _path_hashes(*retain_paths)
        with patch.object(app, "mark_difference_pair_processed", side_effect=stale_pair_miss):
            assert loaded_view._copy_selected_row("MINE2A")
        assert _path_hashes(*retain_paths) == stale_before
        assert retain_item.id not in app._difference_processed
        assert current_item.id in app._difference_processed
        assert app.undo_stack and current_item.id in app.undo_stack[-1]["item_ids"]
        loaded_view._undo_last_action()
        assert not app._difference_processed

        # Region scope must apply the same current-pair authority: an old
        # browser selection cannot cause the region's first row to be marked.
        loaded_view._selected_difference_item_id = retain_item.id
        with patch.object(app, "mark_difference_pair_processed", side_effect=stale_pair_miss):
            assert loaded_view._copy_selected_region("MINE2A")
        assert _path_hashes(*retain_paths) == stale_before
        assert retain_item.id not in app._difference_processed
        assert current_item.id in app._difference_processed
        loaded_view._undo_last_action()
        assert not app._difference_processed

        # Legacy callers can present a cache whose pair ID no longer matches
        # the current browser snapshot.  The production row fallback must
        # still return the real IDs so a retain action remains undoable.
        loaded_view.selected_pair_idx = retain_pair
        loaded_view._last_selected_line = loaded_view.row_to_line.get(retain_pair, 1)
        loaded_view._selected_difference_item_id = None
        pair_marker = app.mark_difference_pair_processed
        pair_miss = lambda sheet, pair_idx, columns=None, processed=True: pair_marker(
            sheet, int(pair_idx) + 1000000, columns, processed
        )
        pair_fallback_before = _path_hashes(*retain_paths)
        with patch.object(app, "mark_difference_pair_processed", side_effect=pair_miss):
            assert loaded_view._copy_selected_row("MINE2A")
        assert _path_hashes(*retain_paths) == pair_fallback_before
        assert retain_item.id in app._difference_processed
        assert app.undo_stack and app.undo_stack[-1]["kind"] == "difference_state"
        loaded_view._undo_last_action()
        assert retain_item.id not in app._difference_processed

        # The same pair miss must flow through the region accumulator without
        # duplicate IDs; one undo restores every marked item and never writes.
        projection_generation = loaded_view._column_projection_generation
        with patch.object(app, "mark_difference_pair_processed", side_effect=pair_miss):
            assert loaded_view._copy_selected_region("MINE2A")
        assert _path_hashes(*retain_paths) == pair_fallback_before
        assert loaded_view._column_projection_generation == projection_generation
        assert len(app.undo_stack[-1]["item_ids"]) == len(
            set(app.undo_stack[-1]["item_ids"])
        )
        loaded_view._undo_last_action()
        assert not app._difference_processed

        # Global retain owns the same real DifferenceItem IDs directly and is
        # intentionally independent of pair lookup; keep the mismatch probe
        # active to prove it does not regress into an unsupported action.
        loaded_view.selected_pair_idx = current_pair
        loaded_view._selected_difference_item_id = retain_item.id
        with patch.object(app, "mark_difference_pair_processed", side_effect=pair_miss):
            assert loaded_view._copy_all_safe_sheet_differences("MINE2A")
        assert _path_hashes(*retain_paths) == pair_fallback_before
        assert loaded_view._column_projection_generation == projection_generation
        assert retain_item.id in app._difference_processed
        assert current_item.id in app._difference_processed
        loaded_view._undo_last_action()
        assert not app._difference_processed
        loaded_view.selected_pair_idx = retain_pair
        loaded_view._last_selected_line = loaded_view.row_to_line.get(retain_pair, 1)
        loaded_view._selected_difference_item_id = retain_item.id

        assert loaded_view._copy_selected_region("MINE2A")
        assert _path_hashes(*retain_paths) == retain_disk_before
        assert loaded_view._column_projection_generation == projection_generation
        assert retain_item.id in app._difference_processed
        loaded_view._undo_last_action()
        assert retain_item.id not in app._difference_processed
        assert loaded_view._copy_all_safe_sheet_differences("MINE2A")
        assert _path_hashes(*retain_paths) == retain_disk_before
        assert loaded_view._column_projection_generation == projection_generation
        assert loaded_view._undo_last_action() is None
        assert not app._difference_processed

        # Double-click is a browse/details gesture. Guard the mutation entry
        # point so an accidental implementation regression fails loudly.
        class _DoubleClick:
            x = 1
            y = 1

        double_click_before = _path_hashes(*retain_paths)
        with patch.object(loaded_view, "_copy_cell", side_effect=AssertionError("double-click wrote")):
            loaded_view._on_double_click_navigate(loaded_view.left, _DoubleClick(), "A")
        app.root.update_idletasks()
        assert _path_hashes(*retain_paths) == double_click_before

        # Production column projection probe: Alpha/Beta are reordered, a
        # Mine-only column carries a numeric value, and a Theirs-side cell is
        # explicitly None. A valid projection's missing side must not borrow
        # the adjacent logical column.
        columns_base = root_dir / "columns-base.xlsx"
        columns_mine = root_dir / "columns-mine.xlsx"
        columns_theirs = root_dir / "columns-theirs.xlsx"
        _write_column_book(columns_base, ["id", "Alpha", "Beta"], [1, "base-alpha", 101])
        _write_column_book(
            columns_mine,
            ["id", "Beta", "Alpha", "MineOnly"],
            [1, 202, "mine-alpha", 0],
        )
        _write_column_book(
            columns_theirs,
            ["id", "Beta", "Alpha", "MineOnly"],
            [1, 203, "theirs-alpha", None],
        )
        columns_app = smt.SowMergeApp(
            str(columns_mine),
            str(columns_theirs),
            merge_mode=True,
            base_path=str(columns_base),
            raw_base=str(columns_base),
            raw_mine=str(columns_mine),
            raw_theirs=str(columns_theirs),
            launch_context=smt.MergeLaunchContext(smt.MergeScenario.UPDATE_CONFLICT),
        )
        columns_view = _wait_for_view(columns_app, "Columns")
        assert columns_view is not None
        column_items = tuple(item for item in columns_app.get_difference_items(5000) if item.sheet == "Columns")
        slots = columns_view.column_projection.model.slots
        assert any(slot.base_col is None for slot in slots), slots
        for item in column_items:
            if not item.column:
                continue
            slot = columns_view.column_projection.slot(item.column)
            if slot is not None and slot.base_col is None:
                assert item.base_value is None, item
        assert any(item.mine_value in {"0", 0} for item in column_items), column_items

        # The same retain action uses Source/Target vocabulary in a real
        # cross-branch launch, while still remaining a no-write acknowledgement.
        cross_context = smt.build_merge_launch_context(
            str(base) + ".merge-left.r10",
            str(mine),
            str(theirs) + ".merge-right.r11",
        )
        cross_app = smt.SowMergeApp(
            str(mine),
            str(theirs),
            merge_mode=True,
            base_path=str(base),
            raw_base=cross_context.source_base_path,
            raw_mine=cross_context.mine_path,
            raw_theirs=cross_context.theirs_path,
            launch_context=cross_context,
        )
        cross_items = _wait_for_items(cross_app, "Data")
        cross_item = next(item for item in cross_items if item.column)
        cross_view = cross_app.sheet_views["Data"]
        cross_view.selected_pair_idx = int(
            cross_item.id.split(":pair:", 1)[1].split(":", 1)[0]
        )
        cross_view._last_selected_line = cross_view.row_to_line.get(
            cross_view.selected_pair_idx, 1
        )
        cross_view._selected_difference_item_id = cross_item.id
        assert cross_view._format_direction_action("MINE2A") == "保留 Target Working 行"
        cross_paths = (mine, theirs, cross_app.file_a, cross_app.file_b)
        if not cross_app._edit_workbooks_ready():
            assert cross_view._copy_selected_row("MINE2A") is False
        assert _wait_for_edit_ready(cross_app)
        cross_mine_before = _path_hashes(*cross_paths)
        assert cross_view._copy_selected_row("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_item.id in cross_app._difference_processed
        cross_view._undo_last_action()
        assert not cross_app._difference_processed

        cross_projection_generation = cross_view._column_projection_generation
        cross_pair = cross_view.selected_pair_idx
        cross_current_item = next(
            item
            for item in cross_items
            if item.column
            and ":pair:" in item.id
            and item.id.split(":pair:", 1)[1].split(":", 1)[0] != str(cross_pair)
        )
        cross_current_pair = int(
            cross_current_item.id.split(":pair:", 1)[1].split(":", 1)[0]
        )
        cross_view.selected_pair_idx = cross_current_pair
        cross_view._last_selected_line = cross_view.row_to_line.get(cross_current_pair, 1)
        cross_view._selected_difference_item_id = cross_item.id
        cross_stale_marker = cross_app.mark_difference_pair_processed
        cross_stale_miss = lambda sheet, pair_idx, columns=None, processed=True: cross_stale_marker(
            sheet, int(pair_idx) + 1000000, columns, processed
        )
        with patch.object(cross_app, "mark_difference_pair_processed", side_effect=cross_stale_miss):
            assert cross_view._copy_selected_row("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_item.id not in cross_app._difference_processed
        assert cross_current_item.id in cross_app._difference_processed
        assert cross_app.undo_stack and cross_current_item.id in cross_app.undo_stack[-1]["item_ids"]
        cross_view._undo_last_action()
        assert not cross_app._difference_processed

        with patch.object(cross_app, "mark_difference_pair_processed", side_effect=cross_stale_miss):
            assert cross_view._copy_selected_region("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_item.id not in cross_app._difference_processed
        assert cross_current_item.id in cross_app._difference_processed
        assert cross_view._column_projection_generation == cross_projection_generation
        cross_view._undo_last_action()
        assert not cross_app._difference_processed

        with patch.object(cross_app, "mark_difference_pair_processed", side_effect=cross_stale_miss):
            assert cross_view._copy_all_safe_sheet_differences("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_item.id in cross_app._difference_processed
        assert cross_current_item.id in cross_app._difference_processed
        cross_view._undo_last_action()
        assert not cross_app._difference_processed

        cross_view.selected_pair_idx = cross_pair
        cross_view._last_selected_line = cross_view.row_to_line.get(cross_pair, 1)
        # Exercise the pair-ID compatibility fallback through the real
        # cross-branch entry points as well.  A deliberately stale pair lookup
        # must still produce one deduplicated, undoable DifferenceItem action.
        cross_view._selected_difference_item_id = None
        cross_pair_marker = cross_app.mark_difference_pair_processed
        cross_pair_miss = lambda sheet, pair_idx, columns=None, processed=True: cross_pair_marker(
            sheet, int(pair_idx) + 1000000, columns, processed
        )
        cross_projection_generation = cross_view._column_projection_generation
        with patch.object(cross_app, "mark_difference_pair_processed", side_effect=cross_pair_miss):
            assert cross_view._copy_selected_row("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_item.id in cross_app._difference_processed
        assert cross_app.undo_stack and cross_app.undo_stack[-1]["kind"] == "difference_state"
        cross_view._undo_last_action()
        assert not cross_app._difference_processed

        with patch.object(cross_app, "mark_difference_pair_processed", side_effect=cross_pair_miss):
            assert cross_view._copy_selected_region("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_view._column_projection_generation == cross_projection_generation
        assert len(cross_app.undo_stack[-1]["item_ids"]) == len(
            set(cross_app.undo_stack[-1]["item_ids"])
        )
        cross_view._undo_last_action()
        assert not cross_app._difference_processed

        with patch.object(cross_app, "mark_difference_pair_processed", side_effect=cross_pair_miss):
            assert cross_view._copy_all_safe_sheet_differences("MINE2A")
        assert _path_hashes(*cross_paths) == cross_mine_before
        assert cross_view._column_projection_generation == cross_projection_generation
        cross_view._undo_last_action()
        assert not cross_app._difference_processed
        app.root.update_idletasks()
        loaded_view._keep_panes_equal()
        loaded_view._keep_panes_equal()
        pane_key = loaded_view._pane_settings_key
        pane_saved = app.settings.get("pane_sashes", {}).get(pane_key, {})
        assert ("three" if loaded_view._is_three_way_expanded() else "two") in pane_saved
        loaded_view.three_way_var.set(0)
        loaded_view._toggle_three_way_view()
        app.root.update_idletasks()
        loaded_view._keep_panes_equal()
        loaded_view._keep_panes_equal()
        assert "two" in app.settings.get("pane_sashes", {}).get(pane_key, {})
        loaded_view.three_way_var.set(1)
        loaded_view._toggle_three_way_view()
        app.root.update_idletasks()
        loaded_view._keep_panes_equal()
        loaded_view._keep_panes_equal()
        assert "three" in app.settings.get("pane_sashes", {}).get(pane_key, {})

        # Vertical module layout uses the same stable pane key and remains
        # bounded when a grip is dragged beyond the window budget.
        vertical_before = loaded_view._vertical_layout.lower_height

        class _GripEvent:
            def __init__(self, y_root):
                self.y_root = y_root

        loaded_view._on_main_vertical_grip_press(_GripEvent(500))
        loaded_view._on_main_vertical_grip_motion(_GripEvent(100))
        vertical_saved = app.settings.get("vertical_sashes", {}).get(pane_key, {})
        assert loaded_view._vertical_layout.lower_height >= vertical_before
        assert vertical_saved.get("lower_height") == loaded_view._vertical_layout.lower_height
        loaded_view._on_hover_vertical_grip_press(_GripEvent(500))
        loaded_view._on_hover_vertical_grip_motion(_GripEvent(100))
        assert loaded_view._vertical_layout.hover_height >= 48
        app._sheet_nav_height = 120
        app._apply_sheet_nav_height()
        assert app.settings.get("sheet_nav_height") == app._sheet_nav_height
        app._reset_workspace_layout()
        assert loaded_view._vertical_layout == VerticalLayout.default(
            three_way=loaded_view._is_three_way_enabled()
        ) or loaded_view._vertical_layout.lower_height >= 132

        # Completing a progress owner invalidates its generation and drops a
        # replaceable broker payload, so a late callback cannot reopen it.
        class _ProgressView:
            sheet = "Data"
            _only_diff_async_build_seq = 7

        progress_view = _ProgressView()
        app._only_diff_progress_owner = (progress_view, 7)
        app._exact_broker_pending = (progress_view, 7, lambda: None)
        assert app._finish_only_diff_progress(progress_view, 7, outcome="success")
        assert app._only_diff_progress_owner is None
        assert app._exact_broker_pending is None
        assert progress_view._only_diff_async_build_seq == 8

        # The unopened Sheet must be projected from the retained cache without
        # constructing a second SheetView or rescanning its workbook.
        cached_items = _wait_for_items(app, "Cached")
        _assert_common_row_projection(cached_items)
        assert app.sheet_views.get("Cached") is None
        assert "Cached" in app._sheet_cache_store

        # Real two-way VCS role/value semantics: A=Base, B=Mine, no Theirs.
        two_base = root_dir / "vcs-base.xlsx"
        two_mine = root_dir / "vcs-mine.xlsx"
        _write_book(two_base, ["Base text", "base-2"])
        _write_book(two_mine, ["Mine text", "mine-2"])
        two_way = smt.SowMergeApp(
            str(two_base),
            str(two_mine),
            raw_base=str(two_base) + ".r10",
            raw_mine=str(two_mine),
        )
        assert Path(two_way.file_b).resolve() == two_mine.resolve()
        assert Path(two_way.file_a).resolve() != two_base.resolve()
        two_items = _wait_for_items(two_way, "Data")
        assert two_items and all(
            item.theirs_value is None
            and (item.source_side, item.target_side) == ("base", "mine")
            for item in two_items
            if item.column
        )
        two_view = two_way.sheet_views["Data"]
        assert getattr(two_way._wb_a_val, "read_only", False)
        assert getattr(two_way._wb_b_val, "read_only", False)
        if not two_way._edit_workbooks_ready():
            two_view.selected_pair_idx = 0
            assert two_view._copy_selected_row("A2B") is False
        assert _wait_for_edit_ready(two_way)
        two_view.selected_pair_idx = 0
        base_before = _digest(two_base)
        mine_before = _digest(two_mine)
        assert two_view._copy_selected_row("A2B")
        with patch.object(two_way, "_confirm_overwrite", return_value=True), patch.object(
            smt.messagebox, "showinfo", return_value=None
        ):
            two_way.save_b_inplace()
        assert _digest(two_base) == base_before
        assert _digest(two_mine) != mine_before
        assert _digest(Path(two_way.file_a)) == base_before
        assert _digest(Path(two_way.file_b)) != mine_before
        saved = load_workbook(two_mine)
        try:
            assert saved.active["A1"].value == "Base text"
        finally:
            saved.close()
        assert two_view._guard_copy_direction("B2A", "测试反向覆盖") is False
        print("SMOKE_DIFFERENCE_BROWSER_REAL_WORKBOOKS_OK")
    finally:
        for instance in (two_way, columns_app, cross_app, app):
            if instance is not None:
                try:
                    instance._shutdown_root()
                except Exception:
                    pass
        shutil.rmtree(root_dir, ignore_errors=True)


if __name__ == "__main__":
    main()
