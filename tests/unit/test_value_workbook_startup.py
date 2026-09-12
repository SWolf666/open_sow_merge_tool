"""Headless coverage for the read-only startup value-workbook gate."""

from __future__ import annotations

import threading

from openpyxl import Workbook, load_workbook

from sow_merge_tool import legacy_core as smt


def _write_book(path, value: str) -> None:
    workbook = Workbook()
    workbook.active["A1"] = value
    workbook.save(path)
    workbook.close()


def _value_gate_stub(app, left, right) -> None:
    app.has_base = False
    app._file_a_val_path = str(left)
    app._file_b_val_path = str(right)
    app._wb_a_val = load_workbook(left, data_only=True, read_only=True)
    app._wb_b_val = load_workbook(right, data_only=True, read_only=True)
    app._wb_base_val = None
    app._value_workbooks_regular = False
    app._value_workbook_upgrade_lock = threading.Lock()


def test_startup_values_are_read_only_and_upgrade_before_edit(tmp_path):
    left = tmp_path / "left.xlsx"
    right = tmp_path / "right.xlsx"
    _write_book(left, "left")
    _write_book(right, "right")

    app = object.__new__(smt.SowMergeApp)
    _value_gate_stub(app, left, right)
    assert app._wb_a_val.read_only is True
    assert app._wb_b_val.read_only is True
    assert app._wb_a_val["Sheet"]["A1"].value == "left"
    assert app._wb_b_val["Sheet"]["A1"].value == "right"
    assert not hasattr(app._wb_a_val["Sheet"], "_cells")

    app._ensure_value_workbooks_regular()

    assert app._value_workbooks_regular is True
    assert app._wb_a_val.read_only is False
    assert app._wb_b_val.read_only is False
    assert app._wb_a_val["Sheet"]["A1"].value == "left"
    assert app._wb_b_val["Sheet"]["A1"].value == "right"
    assert hasattr(app._wb_a_val["Sheet"], "_cells")
    app._wb_a_val.close()
    app._wb_b_val.close()


def test_recalc_reload_derives_regular_marker_from_replaced_objects(tmp_path):
    left = tmp_path / "left.xlsx"
    right = tmp_path / "right.xlsx"
    replacement = tmp_path / "replacement.xlsx"
    _write_book(left, "left")
    _write_book(right, "right")
    _write_book(replacement, "replacement")

    app = object.__new__(smt.SowMergeApp)
    _value_gate_stub(app, left, right)
    app._ensure_value_workbooks_regular()
    app._refresh_current_view_after_val_reload = lambda: None
    app._apply_recalc_results(new_a=str(replacement))

    assert app._value_workbooks_regular is True
    assert app._wb_a_val.read_only is False
    assert app._wb_a_val["Sheet"]["A1"].value == "replacement"
    app._wb_a_val.close()
    app._wb_b_val.close()

